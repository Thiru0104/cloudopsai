from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

# Define the data
headers = ["Task ID", "Task Description", "Status", "ETA", "Remarks"]
data = [
    ("1.0", "Prerequisites & Preparation", "", "", ""),
    ("1.1", "Azure Subscription Setup", "Pending", "1 Hour", "Ensure Owner/Contributor access to create resources."),
    ("1.2", "Resource Group Creation", "Pending", "15 Mins", "Create RG (e.g., rg-nsg-tool-prod) in the target region (e.g., East US)."),
    ("1.3", "Azure CLI & Docker Setup", "Pending", "30 Mins", "Install Azure CLI and Docker locally for build/push operations."),
    
    ("2.0", "Infrastructure Provisioning", "", "", ""),
    ("2.1", "Azure Container Registry (ACR)", "Pending", "20 Mins", "Create ACR to host the backend Docker image."),
    ("2.2", "PostgreSQL Database", "Pending", "45 Mins", "Deploy Azure Database for PostgreSQL (Flexible Server). Save credentials."),
    ("2.3", "Redis Cache", "Pending", "30 Mins", "Deploy Azure Cache for Redis (Basic C0 is sufficient for dev/test)."),
    ("2.4", "Azure Key Vault", "Pending", "20 Mins", "Create Key Vault for storing secrets (DB password, API keys)."),
    ("2.5", "Storage Account", "Pending", "15 Mins", "Create Storage Account for logs/reports (cloudopsai-reports container)."),
    
    ("3.0", "Backend Deployment (App Service)", "", "", ""),
    ("3.1", "Build & Push Docker Image", "Pending", "30 Mins", "Build backend/Dockerfile and push to ACR (v1.0.0)."),
    ("3.2", "Create App Service Plan", "Pending", "15 Mins", "Create Linux App Service Plan (B1 or P1v2 recommended)."),
    ("3.3", "Create Web App for Containers", "Pending", "20 Mins", "Create Web App linked to the ACR image. Enable Managed Identity."),
    ("3.4", "Configure App Settings", "Pending", "30 Mins", "Add Env Vars: DATABASE_URL, REDIS_URL, AZURE_OPENAI_KEY, etc."),
    ("3.5", "Database Migrations", "Pending", "15 Mins", "Run Alembic migrations to initialize the production DB schema."),
    
    ("4.0", "Frontend Deployment (Static Web App)", "", "", ""),
    ("4.1", "Build Frontend", "Pending", "15 Mins", "Run npm run build to generate the dist folder."),
    ("4.2", "Create Static Web App", "Pending", "20 Mins", "Deploy using Azure Static Web Apps (Free or Standard tier)."),
    ("4.3", "Configure API Link", "Pending", "10 Mins", "Set VITE_API_URL to the Backend Web App URL."),
    ("4.4", "CORS Configuration", "Pending", "10 Mins", "Add the Frontend URL to the Backend's Allowed Origins."),
    
    ("5.0", "AI & Security Integration", "", "", ""),
    ("5.1", "Azure OpenAI Access", "Pending", "30 Mins", "Ensure the deployed backend IP/Identity has access to Azure OpenAI."),
    ("5.2", "Managed Identity Role Assignments", "Pending", "20 Mins", "Grant Web App Key Vault Secrets User and Storage Blob Data Contributor roles."),
    
    ("6.0", "Final Validation", "", "", ""),
    ("6.1", "Health Check", "Pending", "15 Mins", "Verify backend /health and frontend loading."),
    ("6.2", "E2E Testing", "Pending", "1 Hour", "Run a full validation workflow (Validation -> AI Recs -> PDF Download).")
]

# File path (saving to root directory)
file_path = os.path.abspath(os.path.join(os.getcwd(), "../Azure_Deployment_Plan.xlsx"))

def create_excel_file():
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = "Deployment Plan"

        # Define styles
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        center_aligned = Alignment(horizontal="center", vertical="center")
        left_aligned = Alignment(horizontal="left", vertical="center")
        thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

        # Section header style
        section_font = Font(bold=True, color="000000")
        section_fill = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")

        # Write Headers
        ws.append(headers)
        for col_num, cell in enumerate(ws[1], 1):
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_aligned
            cell.border = thin_border

        # Write Data
        for row_data in data:
            ws.append(row_data)
            current_row = ws.max_row
            
            # Check if it's a section header (Task ID ends in .0)
            is_section = row_data[0].endswith(".0")
            
            for col_num, cell in enumerate(ws[current_row], 1):
                cell.border = thin_border
                cell.alignment = left_aligned
                
                if is_section:
                    cell.font = section_font
                    cell.fill = section_fill

        # Adjust Column Widths
        column_widths = {
            'A': 10,  # Task ID
            'B': 40,  # Task Description
            'C': 15,  # Status
            'D': 15,  # ETA
            'E': 60   # Remarks
        }

        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width

        # Save file
        wb.save(file_path)
        print(f"Successfully created {file_path}")

    except Exception as e:
        print(f"Error creating Excel file: {e}")

if __name__ == "__main__":
    create_excel_file()
